# -*- coding: utf-8 -*-
import os
import sys
sys.stdout.reconfigure(encoding='utf-8')

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout_box(doc, text_list, box_type="info", title=""):
    """
    box_type: 'info' (xanh dương), 'tip' (xanh lá), 'warning' (vàng cam)
    """
    colors = {
        "info": {"bg": "F0F7FF", "border": "0B57D0", "title_color": RGBColor(11, 87, 208)},
        "tip": {"bg": "F0FDF4", "border": "16A34A", "title_color": RGBColor(22, 163, 74)},
        "warning": {"bg": "FFFBEB", "border": "D97706", "title_color": RGBColor(217, 119, 6)},
        "purple": {"bg": "FAF5FF", "border": "853BCE", "title_color": RGBColor(133, 59, 206)}
    }
    cfg = colors.get(box_type, colors["info"])

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, cfg["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

    # Set left border thick, others none
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{cfg['border']}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    if title:
        run_title = p.add_run(title + "\n")
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(10.5)
        run_title.font.bold = True
        run_title.font.color.rgb = cfg["title_color"]

    for idx, line in enumerate(text_list):
        run = p.add_run(line + ("\n" if idx < len(text_list) - 1 else ""))
        run.font.name = "Segoe UI"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_image_with_caption(doc, img_path, caption, width_inches=5.6):
    if not os.path.exists(img_path):
        print(f"Warning: Image not found: {img_path}")
        return

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run(caption)
    run_cap.font.name = "Segoe UI"
    run_cap.font.size = Pt(9)
    run_cap.font.italic = True
    run_cap.font.color.rgb = RGBColor(100, 116, 139)

def build_docx_guide():
    doc = Document()
    base_dir = r"D:\Projects\logistics-website"
    img_dir = os.path.join(base_dir, "docs", "railway_guide", "images")
    imgDir = img_dir  # alias for backward compatibility
    output_docx = os.path.join(base_dir, "HUONG_DAN_DANG_KY_RAILWAY_BANG_GMAIL.docx")

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    navy_dark = RGBColor(15, 61, 98)      # #0F3D62
    navy_blue = RGBColor(30, 95, 140)     # #1E5F8C
    purple_railway = RGBColor(133, 59, 206) # #853BCE
    text_dark = RGBColor(30, 41, 59)

    # ═════════════════════════════════════════════════════════════════════════
    # HEADER / COVER TITLE
    # ═════════════════════════════════════════════════════════════════════════
    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_badge.paragraph_format.space_before = Pt(0)
    p_badge.paragraph_format.space_after = Pt(4)
    run_badge = p_badge.add_run("TÀI LIỆU HƯỚNG DẪN VẬN HÀNH HẠ TẦNG ĐÁM MÂY — DỰ ÁN LOGISTICS TMS")
    run_badge.font.name = "Segoe UI"
    run_badge.font.size = Pt(9.5)
    run_badge.font.bold = True
    run_badge.font.color.rgb = purple_railway

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("HƯỚNG DẪN ĐĂNG KÝ & SỬ DỤNG\nTÀI KHOẢN RAILWAY BẰNG GMAIL")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = navy_dark

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Cẩm nang từng bước kèm hình ảnh thực tế — Dành cho Chủ doanh nghiệp, Quản lý & Người dùng không chuyên kỹ thuật (Non-Developer)")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(10.5)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Info summary table
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        [("Dự án ứng dụng:", " Logistics TMS (Spider Express)"), ("Nền tảng Cloud:", " Railway.com (Singapore Datacenter)")],
        [("Tỷ giá tham chiếu:", " 1 USD = 26.000 VNĐ"), ("Đối tượng độc giả:", " Ban Giám Đốc, Kế Toán, Quản Lý Vận Hành")],
        [("Cam kết hình ảnh:", " 100% Ảnh chụp thực tế website & Official Docs CDN"), ("Tiêu chuẩn minh bạch:", " Không dùng ảnh mô phỏng/mockup tự chế")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (k, v) in enumerate(row):
            c = meta_table.cell(r_idx, c_idx)
            set_cell_background(c, "F8FAFC")
            set_cell_margins(c, top=60, bottom=60, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r_k = p.add_run(k)
            r_k.font.name = "Segoe UI"; r_k.font.size = Pt(8.5); r_k.font.bold = True; r_k.font.color.rgb = navy_blue
            r_v = p.add_run(v)
            r_v.font.name = "Segoe UI"; r_v.font.size = Pt(8.5); r_v.font.color.rgb = text_dark

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # PHẦN 1: TỔNG QUAN
    # ═════════════════════════════════════════════════════════════════════════
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    r = h1.add_run("1. TỔNG QUAN: RAILWAY LÀ GÌ VÀ VÌ SAO CHỌN GMAIL?")
    r.font.name = "Segoe UI"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = navy_dark

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(
        "Đối với doanh nghiệp, Railway có thể được hiểu đơn giản như một «Mặt bằng công nghệ cao trên đám mây». "
        "Nơi đây lưu trữ, vận hành liên tục 24/7 toàn bộ phần mềm Quản lý Vận tải & Kho bãi Logistics TMS, được phân chia thành 3 Cụm độc lập: "
        "(1) Cụm Vận hành Thật Production, (2) Cụm Thử nghiệm Phát triển Dev, và (3) Cụm Tự động hóa Workflow & Cloud AI Worker độc lập kèm cơ sở dữ liệu riêng và ổ đĩa lưu trữ mã nguồn chuyên biệt."
    )
    run.font.name = "Segoe UI"; run.font.size = Pt(10); run.font.color.rgb = text_dark

    add_callout_box(
        doc,
        [
            "• Không cần ghi nhớ mật khẩu: Bạn không cần phải đau đầu đặt mật khẩu có chữ hoa, số hay ký tự đặc biệt.",
            "• Cơ chế Magic Link & Google SSO: Đăng nhập an toàn tuyệt đối thông qua tài khoản Google hoặc liên kết gửi thẳng vào hộp thư Gmail.",
            "• Tránh hoàn toàn nguy cơ bị lộ mật khẩu: Hacker không thể dò mật khẩu vì hệ thống không hề lưu mật khẩu của bạn.",
            "• Quản lý tập trung: Toàn bộ hóa đơn, thông báo hệ thống và báo cáo tài nguyên đều gửi về Gmail doanh nghiệp."
        ],
        box_type="tip",
        title="🌟 Ưu điểm vượt trội khi sử dụng Gmail để đăng ký Railway:"
    )

    # ═════════════════════════════════════════════════════════════════════════
    # PHẦN 2: CHUẨN BỊ
    # ═════════════════════════════════════════════════════════════════════════
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    r = h2.add_run("2. CHUẨN BỊ TRƯỚC KHI BẮT ĐẦU (DÀNH CHO NGƯỜI MỚI)")
    r.font.name = "Segoe UI"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = navy_dark

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Để quá trình diễn ra suôn sẻ chỉ trong 1 - 2 phút, bạn chỉ cần chuẩn bị sẵn:")
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    prep_items = [
        ("1. Địa chỉ Gmail hoạt động:", " Chuẩn bị địa chỉ Gmail cá nhân hoặc Gmail công ty (Ví dụ: doanhnghiep.tms@gmail.com). Đảm bảo bạn đang đăng nhập sẵn hộp thư này trên trình duyệt."),
        ("2. Trình duyệt web:", " Khuyên dùng Google Chrome, Microsoft Edge, Cốc Cốc hoặc Safari trên máy tính hoặc laptop."),
        ("3. Thẻ thanh toán quốc tế (Tùy chọn):", " Thẻ Visa hoặc Mastercard (Debit hoặc Credit). Chỉ cần chuẩn bị nếu bạn là Chủ doanh nghiệp hoặc Kế toán chịu trách nhiệm thanh toán Gói Pro ($20/tháng ≈ 520.000 đ/tháng). Nếu bạn chỉ được mời vào xem hệ thống, bạn hoàn toàn MIỄN PHÍ và KHÔNG CẦN thẻ ngân hàng.")
    ]
    for bold_txt, normal_txt in prep_items:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(1)
        p_item.paragraph_format.space_after = Pt(3)
        p_item.paragraph_format.left_indent = Inches(0.2)
        r_b = p_item.add_run(bold_txt)
        r_b.font.name = "Segoe UI"; r_b.font.size = Pt(9.5); r_b.font.bold = True; r_b.font.color.rgb = navy_blue
        r_n = p_item.add_run(normal_txt)
        r_n.font.name = "Segoe UI"; r_n.font.size = Pt(9.5); r_n.font.color.rgb = text_dark

    # ═════════════════════════════════════════════════════════════════════════
    # PHẦN 3: HƯỚNG DẪN CHI TIẾT TỪNG BƯỚC
    # ═════════════════════════════════════════════════════════════════════════
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True
    r = h3.add_run("3. CÁC BƯỚC ĐĂNG KÝ TÀI KHOẢN CHI TIẾT (CÓ HÌNH ẢNH MINH HỌA)")
    r.font.name = "Segoe UI"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = navy_dark

    # Bước 1
    p_step1 = doc.add_paragraph()
    p_step1.paragraph_format.space_before = Pt(6)
    p_step1.paragraph_format.space_after = Pt(2)
    p_step1.paragraph_format.keep_with_next = True
    r = p_step1.add_run("Bước 1: Truy cập trang đăng nhập chính thức của Railway")
    r.font.name = "Segoe UI"; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = purple_railway

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("• Bạn mở trình duyệt web và nhập đường dẫn sau: ")
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark
    r_url = p.add_run("https://railway.com/login\n")
    r_url.font.name = "Segoe UI"; r_url.font.size = Pt(10); r_url.font.bold = True; r_url.font.color.rgb = purple_railway
    r2 = p.add_run("• Tại màn hình xuất hiện, nút lớn màu tím «Continue with GitHub» là dành riêng cho dân lập trình viên. ")
    r2.font.name = "Segoe UI"; r2.font.size = Pt(10); r2.font.color.rgb = text_dark
    r3 = p.add_run("Đối với người dùng thông thường, bạn hãy nhấp chuột vào dòng chữ nhỏ bên dưới: ")
    r3.font.name = "Segoe UI"; r3.font.size = Pt(10); r3.font.color.rgb = text_dark
    r4 = p.add_run("«Log in using email»")
    r4.font.name = "Segoe UI"; r4.font.size = Pt(10); r4.font.bold = True; r4.font.color.rgb = RGBColor(11, 87, 208)

    add_image_with_caption(
        doc,
        os.path.join(img_dir, "real_step1_cropped.png"),
        "Hình 1: Ảnh chụp thực tế tại railway.com/login — Bấm vào dòng chữ 'Log in using email' bên dưới nút tím GitHub",
        width_inches=4.6
    )

    # Bước 2
    p_step2 = doc.add_paragraph()
    p_step2.paragraph_format.space_before = Pt(8)
    p_step2.paragraph_format.space_after = Pt(2)
    p_step2.paragraph_format.keep_with_next = True
    r = p_step2.add_run("Bước 2: Lựa chọn 1 trong 2 cách đăng ký bằng Gmail")
    r.font.name = "Segoe UI"; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = purple_railway

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Sau khi bấm «Log in using email», giao diện sẽ mở ra khung xác thực chuyên dụng cho Email. Tại đây bạn có 2 sự lựa chọn cực kỳ thuận tiện:")
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    add_image_with_caption(
        doc,
        os.path.join(imgDir, "real_step2_cropped.png"),
        "Hình 2: Ảnh chụp thực tế form đăng nhập Railway — Lựa chọn 'Continue with Google' (trên) hoặc điền địa chỉ Gmail (dưới)",
        width_inches=4.6
    )

    # So sánh 2 cách
    comp_tbl = doc.add_table(rows=3, cols=3)
    comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Phương thức", "Cách thức thực hiện", "Đánh giá cho người mới"]
    for i, h in enumerate(headers):
        c = comp_tbl.cell(0, i)
        set_cell_background(c, "1E5F8C")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Segoe UI"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)

    rows_data = [
        (
            "CÁCH 1 (Khuyên dùng):\nNút «Continue with Google»",
            "Bấm trực tiếp vào nút có biểu tượng chữ G nhiều màu. Trình duyệt hiện màn hình chọn tài khoản Gmail, bạn chỉ việc bấm chọn là vào ngay lập tức.",
            "⭐ NHANH NHẤT (1 Click)\nKhông cần mở hòm thư Gmail, không cần đợi link xác thực."
        ),
        (
            "CÁCH 2:\nNhập Gmail vào ô trống & bấm «Continue with Email»",
            "Gõ địa chỉ Gmail của bạn vào ô 'hello@email.com'. Đánh dấu xác thực 'Verify you are human' rồi bấm 'Continue with Email'. Railway sẽ gửi 1 đường link đăng nhập về hộp thư.",
            "Phù hợp nếu bạn muốn kiểm soát qua hòm thư hoặc dùng Gmail khác với tài khoản đang đăng nhập trên trình duyệt."
        )
    ]
    for r_idx, r_data in enumerate(rows_data, 1):
        for c_idx, val in enumerate(r_data):
            c = comp_tbl.cell(r_idx, c_idx)
            bg = "F0FDF4" if r_idx == 1 else "F8FAFC"
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Segoe UI"; r.font.size = Pt(8.5); r.font.color.rgb = text_dark

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Minh họa Cách 1
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.space_before = Pt(6)
    p_c1.paragraph_format.space_after = Pt(2)
    p_c1.paragraph_format.keep_with_next = True
    r = p_c1.add_run("▶ Nếu bạn chọn CÁCH 1 («Continue with Google»):")
    r.font.name = "Segoe UI"; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = navy_blue

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "Khi bấm vào nút «Continue with Google», một cửa sổ chọn tài khoản Google quen thuộc sẽ hiện ra. "
        "Bạn chỉ cần nhấp chọn địa chỉ Gmail của mình đang đăng nhập trên máy tính. "
        "Google sẽ bảo mật xác thực danh tính và hệ thống Railway tự động kết nối tài khoản ngay lập tức (1-Click, không cần ghi nhớ mật khẩu)!"
    )
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    # Minh họa Cách 2
    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.space_before = Pt(6)
    p_c2.paragraph_format.space_after = Pt(2)
    p_c2.paragraph_format.keep_with_next = True
    r = p_c2.add_run("▶ Nếu bạn chọn CÁCH 2 («Continue with Email»):")
    r.font.name = "Segoe UI"; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = navy_blue

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Bạn điền địa chỉ Gmail vào ô trống (ví dụ: doanhnghiep.tms@gmail.com). Tích chọn ô xác thực bảo mật người thật «Verify you are human» rồi bấm nút «Continue with Email».")
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    add_image_with_caption(
        doc,
        os.path.join(imgDir, "real_step3_cropped.png"),
        "Hình 3: Ảnh chụp thực tế thao tác nhập Gmail & xác thực bảo mật 'Verify you are human' trước khi nhận link",
        width_inches=4.6
    )

    # Bước 3
    p_step3 = doc.add_paragraph()
    p_step3.paragraph_format.space_before = Pt(8)
    p_step3.paragraph_format.space_after = Pt(2)
    p_step3.paragraph_format.keep_with_next = True
    r = p_step3.add_run("Bước 3: Mở Gmail và nhấp vào liên kết Magic Link (Dành cho Cách 2)")
    r.font.name = "Segoe UI"; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = purple_railway

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "1. Bạn mở một thẻ (tab) mới trên trình duyệt và truy cập vào hộp thư Gmail cá nhân hoặc công ty (mail.google.com).\n"
        "2. Tìm thư mới nhất được gửi từ «Railway Team» (tiêu đề xác thực đăng nhập vào Railway).\n"
        "3. Mở email và bấm vào nút lớn: «Log In to Railway» (hoặc đường liên kết bảo mật đính kèm trong thư).\n"
        "4. Ngay khi bấm vào liên kết, trình duyệt sẽ tự động kích hoạt tài khoản và chuyển bạn vào hệ thống mà không cần bất kỳ mật khẩu nào."
    )
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    add_callout_box(
        doc,
        [
            "• Kiểm tra hộp thư Rác (Spam): Nếu sau 30 giây không thấy email ở Hộp thư đến (Inbox), bạn hãy kiểm tra mục Spam hoặc Thư quảng cáo (Promotions).",
            "• Thời hạn liên kết: Đường dẫn Magic Link có hiệu lực trong vòng 15 phút. Nếu quá hạn, bạn chỉ cần thực hiện lại Bước 2 để nhận link mới.",
            "• Bảo mật tuyệt đối: Chỉ người có quyền truy cập vào hộp thư Gmail mới mở được liên kết này."
        ],
        box_type="warning",
        title="⚠️ Mẹo nhỏ khi kiểm tra hộp thư Gmail:"
    )

    # Bước 4
    p_step4 = doc.add_paragraph()
    p_step4.paragraph_format.space_before = Pt(8)
    p_step4.paragraph_format.space_after = Pt(2)
    p_step4.paragraph_format.keep_with_next = True
    r = p_step4.add_run("Bước 4: Hoàn tất khởi tạo & truy cập Bảng điều khiển (Project Canvas)")
    r.font.name = "Segoe UI"; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = purple_railway

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(
        "Sau khi bấm vào link xác thực hoặc chọn Google, trình duyệt sẽ đưa bạn vào trang quản trị chính thức của Railway (Railway Project Canvas). "
        "Tại đây, bạn sẽ thấy Bảng điều khiển trực quan hiển thị toàn bộ các khối vi dịch vụ (Services) và Cơ sở dữ liệu (PostgreSQL Database) đang chạy ổn định 24/7."
    )
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    add_image_with_caption(
        doc,
        os.path.join(imgDir, "official_project_canvas.png"),
        "Hình 4: Ảnh chụp thực tế từ hệ thống Railway — Bảng điều khiển Project Canvas hiển thị các dịch vụ và cơ sở dữ liệu của dự án TMS",
        width_inches=5.6
    )

    # ═════════════════════════════════════════════════════════════════════════
    # PHẦN 4: HƯỚNG DẪN QUẢN TRỊ & MỜI THÀNH VIÊN
    # ═════════════════════════════════════════════════════════════════════════
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)
    h4.paragraph_format.keep_with_next = True
    r = h4.add_run("4. QUẢN LÝ KHÔNG GIAN LÀM VIỆC & MỜI THÀNH VIÊN (TEAM)")
    r.font.name = "Segoe UI"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = navy_dark

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "Một trong những ưu điểm lớn nhất của gói Railway Pro ($20/tháng) so với các nhà cung cấp khác là chính sách "
        "«UNLIMITED SEATS» (Không giới hạn số lượng thành viên, 0đ phụ phí). Doanh nghiệp có thể mời toàn bộ các thành viên "
        "thuộc Ban giám đốc, Kế toán, Quản lý kho, Điều độ và Đội ngũ lập trình viên cùng vào theo dõi mà không tốn thêm bất kỳ chi phí phát sinh nào."
    )
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    add_image_with_caption(
        doc,
        os.path.join(imgDir, "official_member_view.png"),
        "Hình 5: Ảnh chụp thực tế từ hệ thống Railway — Giao diện quản lý thành viên & phân quyền dự án (Project Members & Roles)",
        width_inches=5.6
    )

    add_image_with_caption(
        doc,
        os.path.join(imgDir, "official_trusted_domains.png"),
        "Hình 6: Ảnh chụp thực tế từ hệ thống Railway — Giao diện Workspace Settings & Cấu hình Tên miền tin cậy (Trusted Domains)",
        width_inches=5.6
    )

    # Bảng phân quyền
    role_tbl = doc.add_table(rows=4, cols=3)
    role_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    role_headers = ["Cấp độ vai trò", "Quyền hạn trong hệ thống", "Đối tượng phù hợp trong doanh nghiệp"]
    for i, h in enumerate(role_headers):
        c = role_tbl.cell(0, i)
        set_cell_background(c, "0F3D62")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Segoe UI"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)

    role_data = [
        ("Owner / Admin\n(Toàn quyền Quản trị)", "Toàn quyền quản trị hệ thống, thêm/bớt thành viên, quản lý thẻ thanh toán và xem toàn bộ nhật ký hoạt động.", "Chủ doanh nghiệp, Giám đốc CNTT, Kế toán trưởng đứng tên thẻ."),
        ("Member\n(Thành viên xem & quản lý)", "Xem trạng thái hoạt động của các dịch vụ, kiểm tra log lỗi, quản lý biến môi trường dự án nhưng KHÔNG được xóa dự án hay đổi thẻ thanh toán.", "Quản lý dự án (PM), Trưởng nhóm kỹ thuật (Tech Lead), QA / Tester."),
        ("Viewer / Deployer\n(Chỉ xem)", "Chỉ được xem trạng thái hoạt động của hệ thống, không được xem mật khẩu hay biến môi trường nhạy cảm.", "Quản lý vận hành kho bãi, nhân viên giám sát ca trực.")
    ]
    for r_idx, r_data in enumerate(role_data, 1):
        for c_idx, val in enumerate(r_data):
            c = role_tbl.cell(r_idx, c_idx)
            set_cell_background(c, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=70, bottom=70, left=90, right=90)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Segoe UI"; r.font.size = Pt(8.5); r.font.color.rgb = text_dark

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ═════════════════════════════════════════════════════════════════════════
    # PHẦN 5: THIẾT LẬP THANH TOÁN (BILLING)
    # ═════════════════════════════════════════════════════════════════════════
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(4)
    h5.paragraph_format.keep_with_next = True
    r = h5.add_run("5. HƯỚNG DẪN THIẾT LẬP THẺ THANH TOÁN & QUẢN LÝ CHI PHÍ")
    r.font.name = "Segoe UI"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = navy_dark

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        "Mục này dành riêng cho Chủ doanh nghiệp hoặc Kế toán chịu trách nhiệm thanh toán cước phí máy chủ hàng tháng. "
        "Railway hỗ trợ thanh toán tự động qua thẻ Visa / Mastercard quốc tế một cách cực kỳ minh bạch và an toàn."
    )
    r.font.name = "Segoe UI"; r.font.size = Pt(10); r.font.color.rgb = text_dark

    bill_steps = [
        ("Bước 1:", " Tại góc trên bên phải màn hình Railway Dashboard, bấm vào ảnh đại diện tài khoản -> Chọn «Workspace Settings» -> Chọn mục «Billing»."),
        ("Bước 2:", " Bấm chọn gói «PRO Plan» ($20/tháng)."),
        ("Bước 3:", " Nhập thông tin thẻ quốc tế của doanh nghiệp: Số thẻ (16 chữ số), Ngày hết hạn (MM/YY), và Mã bảo mật CVC (3 số ở mặt sau thẻ)."),
        ("Bước 4:", " Nhập Tên doanh nghiệp (Company Name) và Địa chỉ thanh toán để hóa đơn xuất ra đúng thông tin doanh nghiệp làm căn cứ kế toán.")
    ]
    for b_s, b_txt in bill_steps:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = Pt(1)
        p_b.paragraph_format.space_after = Pt(3)
        p_b.paragraph_format.left_indent = Inches(0.2)
        r_bs = p_b.add_run(b_s + " ")
        r_bs.font.name = "Segoe UI"; r_bs.font.size = Pt(9.5); r_bs.font.bold = True; r_bs.font.color.rgb = purple_railway
        r_bt = p_b.add_run(b_txt)
        r_bt.font.name = "Segoe UI"; r_bt.font.size = Pt(9.5); r_bt.font.color.rgb = text_dark

    add_image_with_caption(
        doc,
        os.path.join(imgDir, "official_account_settings.png"),
        "Hình 7: Ảnh chụp thực tế từ hệ thống Railway — Giao diện Cài đặt Tài khoản (Account Settings) và liên kết thanh toán",
        width_inches=5.4
    )

    # Bảng phân nhóm 3 cụm hạ tầng
    h5_sub = doc.add_paragraph()
    h5_sub.paragraph_format.space_before = Pt(8)
    h5_sub.paragraph_format.space_after = Pt(2)
    h5_sub.paragraph_format.keep_with_next = True
    r = h5_sub.add_run("Bảng phân bổ chi phí 3 Cụm Hạ tầng chuyên biệt (Áp dụng tỷ giá 1 USD = 26.000 VNĐ):")
    r.font.name = "Segoe UI"; r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = navy_blue

    grp_tbl = doc.add_table(rows=6, cols=4)
    grp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    grp_headers = ["Nhóm Hạ Tầng", "Các thành phần dịch vụ", "Chi phí USD", "Quy đổi VNĐ"]
    for i, h in enumerate(grp_headers):
        c = grp_tbl.cell(0, i)
        set_cell_background(c, "0F3D62")
        set_cell_margins(c, top=80, bottom=80, left=90, right=90)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Segoe UI"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)

    grp_rows = [
        (
            "CỤM 1: Production TMS\n(Vận hành chính thức)",
            "• Frontend Next.js Pro (Xe, Vận đơn, Kế toán)\n• Backend NestJS API Pro (Logic nghiệp vụ)\n• PostgreSQL Pro 16 (DB nghiệp vụ thực tế)\n• Redis Cache Pro (Bộ nhớ đệm tốc độ cao)",
            "$15.00 / thg",
            "390.000 đ"
        ),
        (
            "CỤM 2: Development TMS\n(Môi trường Dev thử nghiệm)",
            "• Frontend Dev (Test giao diện tính năng mới)\n• Backend Dev (Test API, xử lý nghiệp vụ)\n• PostgreSQL Dev (Dữ liệu mẫu thử nghiệm)\n• Redis Cache Dev (Đệm thử nghiệm)",
            "$9.00 / thg",
            "234.000 đ"
        ),
        (
            "CỤM 3: Tự Động Hóa &\nCloud AI Worker\n(Nhóm chuyên biệt)",
            "• n8n Workflow Automation (Zalo ZNS, Webhook)\n• PostgreSQL Automation (DB riêng biệt 100% cho n8n)\n• Telegram Cloud Worker (Nhận lệnh bot & fix bug)\n• Persistent Volume 10GB NVMe (Ổ đĩa chứa Source Code)",
            "$8.30 / thg",
            "215.800 đ"
        ),
        (
            "Cước mạng truyền tải\n(Egress Data Transfer)",
            "Lưu lượng mạng Internet tải hình ảnh/API ra bên ngoài (Dự phòng ~20GB băng thông mở rộng)",
            "$1.00 / thg",
            "26.000 đ"
        ),
        (
            "TỔNG CỘNG HÓA ĐƠN\n(Đã khấu trừ Credit)",
            "Toàn bộ 11 dịch vụ + 1 ổ đĩa SSD 10GB hoạt động liên tục 24/7 (Được tặng sẵn $20 Credit gói Pro)",
            "$33.30 / thg\n(Thực trả: $33.30)",
            "865.800 đ / tháng"
        )
    ]
    for r_idx, r_data in enumerate(grp_rows, 1):
        for c_idx, val in enumerate(r_data):
            c = grp_tbl.cell(r_idx, c_idx)
            is_total = (r_idx == len(grp_rows))
            is_grp3 = (r_idx == 3)
            bg = "EFF6FF" if is_total else ("F5F3FF" if is_grp3 else ("F8FAFC" if r_idx % 2 == 1 else "FFFFFF"))
            set_cell_background(c, bg)
            set_cell_margins(c, top=70, bottom=70, left=80, right=80)
            p = c.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9 if is_total else 8.5)
            r.font.bold = is_total or (c_idx in [0, 2, 3])
            r.font.color.rgb = navy_dark if is_total else (purple_railway if (is_grp3 and c_idx == 0) else text_dark)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_callout_box(
        doc,
        [
            "• Gói Pro tặng sẵn $20 Credit hàng tháng: Tiền gói $20 được trừ thẳng 100% vào tài nguyên máy chủ tiêu thụ trong tháng.",
            "• Cụm 3 (Tự Động Hóa & AI Worker) được bóc tách riêng biệt: Gồm engine n8n ($2.50), database PostgreSQL độc lập ($1.80), Telegram worker daemon ($2.50) và Ổ đĩa gắn ngoài 10GB SSD ($1.50) = $8.30/tháng (215.800 đ).",
            "• Cách ly dữ liệu 100%: n8n và Worker có database riêng biệt ('postgres-automation'), KHÔNG chia sẻ hay dùng chung database với hệ thống TMS cốt lõi, loại trừ hoàn toàn nguy cơ tranh chấp tài nguyên hay xung đột khóa bảng.",
            "• Ổ đĩa lưu trữ mã nguồn chuyên biệt (Persistent Volume 10GB): Gắn trực tiếp vào Cloud Worker để lưu trữ Git submodules, node_modules, build cache và ảnh chụp bằng chứng E2E mà không làm đầy bộ nhớ RAM hay mất dữ liệu khi máy chủ khởi động lại.",
            "• Dự toán chi phí thực tế: Toàn bộ 11 dịch vụ và 1 ổ đĩa SSD tiêu thụ tổng cộng $33.30 USD/tháng ≈ 865.800 VNĐ / tháng (Đã quy đổi theo tỷ giá 26.000 đ = 1$).",
            "• Tải hóa đơn: Hàng tháng Railway tự động gửi Receipt qua Gmail và lưu tại mục Billing để Kế toán tải file PDF về báo cáo."
        ],
        box_type="purple",
        title="💰 Cơ chế tính cước minh bạch của Railway (Áp dụng tỷ giá 26k = 1$):"
    )

    # ═════════════════════════════════════════════════════════════════════════
    # PHẦN 6: FAQ
    # ═════════════════════════════════════════════════════════════════════════
    h6 = doc.add_paragraph()
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(4)
    h6.paragraph_format.keep_with_next = True
    r = h6.add_run("6. NHỮNG CÂU HỎI THƯỜNG GẶP (FAQ) DÀNH CHO NON-DEVELOPER")
    r.font.name = "Segoe UI"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = navy_dark

    faqs = [
        (
            "Câu hỏi 1: Lần sau đăng nhập lại trên máy tính khác hoặc điện thoại thì làm thế nào?",
            "Trả lời: Bạn chỉ cần vào lại trang railway.com/login, bấm 'Continue with Google' (nếu máy có sẵn Gmail) hoặc bấm 'Log in using email' và nhập địa chỉ Gmail. Hệ thống sẽ gửi 1 Magic Link vào Gmail của bạn. Bạn mở thư và bấm link là đăng nhập thành công ngay, hoàn toàn không cần nhập mật khẩu."
        ),
        (
            "Câu hỏi 2: Tôi là Kế toán / Quản lý được mời vào xem hệ thống thì có bị mất phí không?",
            "Trả lời: HOÀN TOÀN KHÔNG MẤT PHÍ. Gói Railway Pro có chính sách Unlimited Seats (Không giới hạn thành viên, 0đ phụ phí). Chi phí duy nhất là cước duy trì hệ thống máy chủ của dự án do Chủ sở hữu thanh toán."
        ),
        (
            "Câu hỏi 3: Thẻ ATM nội địa (Napas) có thanh toán được trên Railway không?",
            "Trả lời: Không. Railway là nền tảng điện toán đám mây quốc tế đặt trụ sở tại Mỹ và máy chủ tại Singapore, do đó bắt buộc sử dụng thẻ thanh toán quốc tế có logo Visa hoặc Mastercard (thẻ Debit hoặc Credit của bất kỳ ngân hàng nào tại Việt Nam như Vietcombank, Techcombank, ACB, MB,...)."
        ),
        (
            "Câu hỏi 4: Nếu tôi vô tình đóng trình duyệt khi đang đăng ký thì có bị sao không?",
            "Trả lời: Không sao cả. Bất cứ lúc nào bạn chỉ cần mở lại Gmail và bấm vào đường link 'Log In to Railway', trình duyệt sẽ tự động đưa bạn về đúng trạng thái đang thao tác."
        ),
        (
            "Câu hỏi 5: Dữ liệu của công ty trên Railway có an toàn và bảo mật không?",
            "Trả lời: Rất an toàn. Máy chủ được đặt tại Datacenter Singapore với tiêu chuẩn bảo mật quốc tế SOC2 và ISO 27001. Hệ thống Logistics TMS của doanh nghiệp còn được thiết lập mạng nội bộ riêng (Wireguard Private Mesh) độc lập 100% giữa môi trường Thử nghiệm (Dev) và Vận hành thật (Production)."
        )
    ]

    for q, a in faqs:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(4)
        p_q.paragraph_format.space_after = Pt(1)
        p_q.paragraph_format.keep_with_next = True
        r_q = p_q.add_run("❓ " + q)
        r_q.font.name = "Segoe UI"; r_q.font.size = Pt(10); r_q.font.bold = True; r_q.font.color.rgb = navy_dark

        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(0)
        p_a.paragraph_format.space_after = Pt(6)
        p_a.paragraph_format.left_indent = Inches(0.2)
        r_a = p_a.add_run("👉 " + a)
        r_a.font.name = "Segoe UI"; r_a.font.size = Pt(9.5); r_a.font.color.rgb = text_dark

    # ═════════════════════════════════════════════════════════════════════════
    # TỔNG KẾT
    # ═════════════════════════════════════════════════════════════════════════
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_callout_box(
        doc,
        [
            "1. Truy cập railway.com/login -> Bấm chọn «Log in using email».",
            "2. Ưu tiên bấm «Continue with Google» (1-Click nhanh nhất) hoặc Nhập Gmail nhận «Magic Link».",
            "3. Bấm xác nhận trong email để vào Bảng điều khiển quản trị (Dashboard).",
            "4. Nhờ Đội ngũ Kỹ thuật thêm tài khoản của bạn vào nhóm quản trị Workspace (Role: Admin / Member)."
        ],
        box_type="tip",
        title="🎯 TÓM TẮT QUY TRÌNH 4 BƯỚC ĐĂNG KÝ SIÊU TỐC CHO NGƯỜI DÙNG MỚI:"
    )

    doc.save(output_docx)
    print(f"🎉 Word document generated successfully at: {output_docx}")

if __name__ == "__main__":
    build_docx_guide()
